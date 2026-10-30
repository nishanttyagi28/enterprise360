"""MIS workbook: sheet-name rules, formatting, README and Excel-safe types (no database)."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reporting.export_mis import MARTS, default_output, sheet_name, write_workbook

pytestmark = pytest.mark.unit


def test_sheet_names_are_safe_unique_and_short() -> None:
    used: set[str] = set()
    a = sheet_name("x" * 40, used)
    b = sheet_name("x" * 40, used)
    assert len(a) == 31 and len(b) <= 31 and a != b
    assert sheet_name("a/b:c*d?[e]", used) == "a_b_c_d__e_"
    assert all(len(m.sheet) <= 31 for m in MARTS)
    assert len({m.sheet.lower() for m in MARTS}) == len(MARTS)


def test_workbook_has_readme_formatted_headers_and_frozen_panes(tmp_path: Path) -> None:
    df = pd.DataFrame(
        {"month": [datetime(2018, 1, 1, tzinfo=UTC)], "gmv": [Decimal("12.50")], "label": ["a"]}
    )
    readme = pd.DataFrame([{"sheet": "Data", "rows": 1}])
    out = tmp_path / "r" / "mis.xlsx"
    names = write_workbook(out, {"Data": df}, readme)
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", names["Data"]]
    ws = wb[names["Data"]]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold and ws["A1"].value == "month"
    assert ws["B2"].value == 12.5  # Decimal -> number
    assert ws["A2"].value == datetime(2018, 1, 1)  # tz-aware -> naive UTC
    assert wb["README"]["A2"].value == "Data"


def test_default_output_is_dated_under_reports() -> None:
    p = default_output(datetime(2026, 10, 31))
    assert p.parent.name == "reports" and p.name == "enterprise360_mis_20261031.xlsx"
