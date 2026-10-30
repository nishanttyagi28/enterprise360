"""``python -m reporting.export_mis [--output PATH] [--max-rows N]``

Exports the key marts to one Excel workbook (default ``reports/enterprise360_mis_YYYYMMDD.xlsx``): one sheet
per mart plus a README sheet (run time, source relation, row counts, notes). Headers are formatted, panes
frozen below the header, autofilter on. Read-only against the warehouse.
"""

from __future__ import annotations

import argparse
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
import psycopg
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from psycopg import sql

from ingestion.config import REPO_ROOT, get_settings
from ingestion.log_config import configure_logging, get_logger
from ingestion.pipeline import connect

log = get_logger(__name__)
EXCEL_MAX_ROWS = 1_048_575  # minus header
_INVALID = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True, slots=True)
class MartSpec:
    schema: str
    table: str
    sheet: str
    note: str
    order_by: str = "1"


MARTS: tuple[MartSpec, ...] = (
    MartSpec(
        "marts",
        "mart_business_overview_monthly",
        "Olist Business Monthly",
        "Olist monthly overview; GMV = item price only (not revenue); use is_reporting_month",
    ),
    MartSpec("marts", "mart_sales_daily", "Olist Sales Daily", "Olist daily sales (BRL)"),
    MartSpec("marts", "mart_sales_by_category_monthly", "Olist Category Monthly", "Olist category x month"),
    MartSpec("marts", "mart_operations_delivery_performance", "Olist Delivery", "Olist delivery performance"),
    MartSpec("marts", "mart_customer_cohorts", "Olist Customer Cohorts", "Olist first-order cohorts"),
    MartSpec(
        "marts",
        "mart_funnel_monthly",
        "Funnel Monthly",
        "Lead cohort (first-contact month) x origin; conversion = sum(converted)/sum(mqls)",
        "1, 4",
    ),
    MartSpec(
        "marts", "mart_funnel_time_to_close", "Funnel Time To Close", "Days-to-close buckets x origin", "2, 3"
    ),
    MartSpec(
        "marts",
        "mart_gmv_forecast",
        "GMV Forecast",
        "Latest forecast run: actual / model / baseline / 95% PI by month",
    ),
    MartSpec("marts", "mart_gmv_anomalies", "GMV Anomalies", "Robust z of residuals; anomaly iff z > 3.5"),
    MartSpec(
        "marts",
        "mart_gmv_forecast_backtest",
        "Forecast Backtest",
        "MAE / sMAPE: model vs baseline",
        "1, 3, 2",
    ),
    MartSpec("uci_marts", "uci_mart_sales_daily", "UCI Sales Daily", "UCI Online Retail II daily (GBP)"),
    MartSpec(
        "uci_marts", "uci_mart_product_monthly", "UCI Product Monthly", "UCI product x month (GBP)", "1, 2"
    ),
    MartSpec(
        "uci_marts", "uci_mart_country_monthly", "UCI Country Monthly", "UCI country x month (GBP)", "1, 2"
    ),
)


def sheet_name(name: str, used: set[str]) -> str:
    """Excel-safe, unique (case-insensitive) sheet name of <= 31 chars."""
    base = _INVALID.sub("_", name).strip("'")[:31] or "Sheet"
    candidate, i = base, 2
    while candidate.lower() in used:
        suffix = f" ({i})"
        candidate, i = base[: 31 - len(suffix)] + suffix, i + 1
    used.add(candidate.lower())
    return candidate


def _excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    """Decimal -> float, tz-aware timestamps -> naive UTC (Excel has no time zones)."""
    out = df.copy()
    for col in out.columns:
        s = out[col]
        if s.dtype == object and s.map(lambda v: isinstance(v, Decimal)).any():
            out[col] = s.map(lambda v: float(v) if isinstance(v, Decimal) else v)
        elif isinstance(s.dtype, pd.DatetimeTZDtype):
            out[col] = s.dt.tz_convert("UTC").dt.tz_localize(None)
        elif s.dtype == object and s.map(lambda v: isinstance(v, datetime) and v.tzinfo is not None).any():
            out[col] = pd.to_datetime(s, utc=True).dt.tz_localize(None)
    return out


def write_workbook(path: Path, frames: Mapping[str, pd.DataFrame], readme: pd.DataFrame) -> dict[str, str]:
    """Write README + one sheet per frame; returns {frame key: sheet name}."""
    path.parent.mkdir(parents=True, exist_ok=True)
    used: set[str] = {"readme"}
    names = {key: sheet_name(key, used) for key in frames}
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        all_frames = {"README": readme, **{names[k]: _excel_safe(v) for k, v in frames.items()}}
        for sheet, df in all_frames.items():
            df.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if len(df.columns):
                ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(len(df), 1) + 1}"
            for idx, col in enumerate(df.columns, start=1):
                sample = df[col].astype(str).head(200).map(len)
                width = max([len(str(col)), *sample.tolist()]) if len(df) else len(str(col))
                ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 60)
    return names


def _relation_exists(conn: psycopg.Connection, spec: MartSpec) -> bool:
    row = conn.execute("select to_regclass(%s) is not null", (f"{spec.schema}.{spec.table}",)).fetchone()
    return bool(row and row[0])


def read_mart(conn: psycopg.Connection, spec: MartSpec, max_rows: int) -> pd.DataFrame:
    query = sql.SQL("select * from {}.{} order by {} limit {}").format(
        sql.Identifier(spec.schema), sql.Identifier(spec.table), sql.SQL(spec.order_by), sql.Literal(max_rows)
    )
    with conn.cursor() as cur:
        cur.execute(query)
        cols = [d.name for d in cur.description or []]
        return pd.DataFrame(cur.fetchall(), columns=cols)


def export(output: Path, *, max_rows: int = EXCEL_MAX_ROWS, specs: Sequence[MartSpec] = MARTS) -> Path:
    settings = get_settings()
    run_at = datetime.now(UTC)
    frames: dict[str, pd.DataFrame] = {}
    meta: list[dict[str, object]] = []
    with connect(settings) as conn:
        for spec in specs:
            relation = f"{spec.schema}.{spec.table}"
            if not _relation_exists(conn, spec):
                meta.append(
                    {
                        "sheet": None,
                        "source_relation": relation,
                        "rows": None,
                        "status": "missing (not built)",
                        "note": spec.note,
                    }
                )
                continue
            df = read_mart(conn, spec, max_rows + 1)
            truncated = len(df) > max_rows
            frames[spec.sheet] = df.head(max_rows)
            meta.append(
                {
                    "sheet": spec.sheet,
                    "source_relation": relation,
                    "rows": min(len(df), max_rows),
                    "status": "truncated" if truncated else "ok",
                    "note": spec.note,
                }
            )
    readme = pd.DataFrame(
        [
            {
                "sheet": "README",
                "source_relation": f"postgres {settings.pg_host}/{settings.pg_db}",
                "rows": None,
                "status": f"generated {run_at:%Y-%m-%d %H:%M:%S} UTC",
                "note": "Enterprise360 MIS export (read-only)",
            },
            *meta,
        ]
    )
    names = write_workbook(output, frames, readme)
    log.info(
        "mis_export_written",
        path=str(output),
        sheets=list(names.values()),
        missing=[m["source_relation"] for m in meta if m["sheet"] is None],
    )
    return output


def default_output(now: datetime | None = None) -> Path:
    return REPO_ROOT / "reports" / f"enterprise360_mis_{(now or datetime.now()):%Y%m%d}.xlsx"


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export key marts to an Excel MIS workbook")
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--max-rows", type=int, default=EXCEL_MAX_ROWS)
    a = parser.parse_args(argv)
    configure_logging(get_settings().log_level)
    path = export(a.output or default_output(), max_rows=min(a.max_rows, EXCEL_MAX_ROWS))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
